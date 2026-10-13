from typing import List, Dict, Any
from pathlib import Path
from pptx import Presentation

def parse_pptx(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extracts text from PowerPoint (.pptx) slides preserving
    the exact 1-indexed slide number.
    """
    prs = Presentation(str(file_path))
    slides_data: List[Dict[str, Any]] = []

    for slide_idx, slide in enumerate(prs.slides):
        slide_number = slide_idx + 1
        slide_texts: List[str] = []

        # Extract title and body text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                # Handle tables inside slides
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        slide_texts.append(" | ".join(row_texts))

        # Also extract speaker notes if available
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                slide_texts.append(f"[Speaker Notes: {notes_text}]")

        combined_text = "\n".join(slide_texts).strip()
        slides_data.append({
            "slide_number": slide_number,
            "locator": f"Slide {slide_number}",
            "text": combined_text or f"[Slide {slide_number} - Visual Content]",
        })

    return slides_data
