from pathlib import Path
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def create_sample_pdf():
    pdf_path = SAMPLE_DIR / "Deep_Learning_Fundamentals.pdf"
    doc = pymupdf.open()

    # Page 1: Introduction & Perceptrons
    p1 = doc.new_page()
    p1.insert_text(
        pymupdf.Point(50, 72),
        "Deep Learning Fundamentals — Chapter 1: The Artificial Neuron\n\n"
        "An Artificial Neuron (or Perceptron) is the fundamental atomic building block of deep neural networks. "
        "Each perceptron computes a weighted sum of its numerical inputs: z = w1*x1 + w2*x2 + ... + wn*xn + b, "
        "where w represents the weight vector and b represents the bias scalar parameter.\n\n"
        "The weighted sum is subsequently passed through a non-linear activation function, f(z), allowing "
        "the neural network to model complex non-linear relationships rather than purely affine transformations. "
        "Historically, the step function and sigmoid function were widely used, but modern architectures "
        "overwhelmingly rely on Rectified Linear Units (ReLU) to avoid vanishing gradients during optimization.",
        fontsize=12,
    )

    # Page 2: Backpropagation & Optimization
    p2 = doc.new_page()
    p2.insert_text(
        pymupdf.Point(50, 72),
        "Deep Learning Fundamentals — Chapter 2: Backpropagation and Gradient Descent\n\n"
        "Backpropagation is an algorithm for calculating the gradient of the loss function with respect to the "
        "weights of the neural network using the mathematical chain rule of calculus.\n\n"
        "1. Forward Pass: Input features propagate forward through network layers to produce the prediction output y_hat.\n"
        "2. Loss Computation: A scalar loss, such as Mean Squared Error (MSE) or Cross-Entropy Loss, quantifies error.\n"
        "3. Backward Pass: Gradients delta_L / delta_w are propagated backwards layer-by-layer.\n"
        "4. Weight Update: Gradient Descent updates the parameters according to w_new = w_old - alpha * grad(w), "
        "where alpha denotes the learning rate hyperparameter. Selecting an excessively large learning rate causes "
        "divergence or oscillation, while an excessively small rate causes prohibitively slow convergence.",
        fontsize=12,
    )

    # Page 3: Activation Functions & Vanishing Gradients
    p3 = doc.new_page()
    p3.insert_text(
        pymupdf.Point(50, 72),
        "Deep Learning Fundamentals — Chapter 3: Activation Functions and Vanishing Gradients\n\n"
        "The Vanishing Gradient problem occurs when gradients shrink exponentially as backpropagation propagates "
        "through many deep layers, preventing earlier layers from updating their weights effectively.\n\n"
        "- Sigmoid Activation: Maps inputs to (0, 1). Its derivative saturates to near-zero for large absolute values (|z| > 4), "
        "which severely triggers vanishing gradients in networks deeper than 3-4 layers.\n"
        "- Tanh Activation: Zero-centered with output range (-1, 1), but still suffers from saturation and vanishing gradients.\n"
        "- ReLU (Rectified Linear Unit): Defined as f(z) = max(0, z). Its derivative is constant (1) for all positive inputs, "
        "substantially mitigating the vanishing gradient issue and enabling the training of very deep neural networks.",
        fontsize=12,
    )

    doc.save(str(pdf_path))
    doc.close()
    print(f"Created sample PDF: {pdf_path}")

def create_sample_pptx():
    pptx_path = SAMPLE_DIR / "Transformer_Architecture_Slides.pptx"
    prs = Presentation()

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Transformer Architecture & Self-Attention"
    slide.placeholders[1].text = "Multimodal AI Hackathon 2026\nModule 4: Attention Mechanisms"

    # Slide 2: Limitations of RNNs
    bullet_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_layout)
    slide2.shapes.title.text = "Why Transformers? Limitations of RNNs and LSTMs"
    tf2 = slide2.placeholders[1].text_frame
    tf2.text = "Traditional sequential models (RNNs/LSTMs) process tokens strictly step-by-step."
    p = tf2.add_paragraph()
    p.text = "Sequential dependency prevents parallel execution across GPU cores during training."
    p = tf2.add_paragraph()
    p.text = "Information loss over long sequences creates a bottleneck for long-distance semantic dependencies."
    p = tf2.add_paragraph()
    p.text = "Transformers replace recurrence entirely with the Attention Mechanism, enabling complete parallelization."

    # Slide 3: Scaled Dot-Product Attention
    slide3 = prs.slides.add_slide(bullet_layout)
    slide3.shapes.title.text = "Scaled Dot-Product Attention Formula"
    tf3 = slide3.placeholders[1].text_frame
    tf3.text = "Attention(Q, K, V) = softmax((Q * K^T) / sqrt(d_k)) * V"
    p = tf3.add_paragraph()
    p.text = "Q (Queries), K (Keys), and V (Values) are linear projections of input embeddings."
    p = tf3.add_paragraph()
    p.text = "Scaling factor 1 / sqrt(d_k) prevents dot products from growing excessively large for large vector dimensions."
    p = tf3.add_paragraph()
    p.text = "Without scaling, the softmax function yields microscopic gradients that cause optimization stagnation."

    # Slide 4: Multi-Head Attention & Positional Encoding
    slide4 = prs.slides.add_slide(bullet_layout)
    slide4.shapes.title.text = "Multi-Head Attention & Positional Encodings"
    tf4 = slide4.placeholders[1].text_frame
    tf4.text = "Multi-Head Attention splits Q, K, and V into h parallel representation subspaces."
    p = tf4.add_paragraph()
    p.text = "This allows the model to jointly attend to information from different representation positions."
    p = tf4.add_paragraph()
    p.text = "Positional Encodings (sinusoidal or learned) are added to input embeddings to inject sequence order."
    p = tf4.add_paragraph()
    p.text = "Because attention operations are permutation-invariant, positional encodings are essential for language syntax."

    prs.save(str(pptx_path))
    print(f"Created sample PPTX: {pptx_path}")

def create_sample_srt():
    srt_path = SAMPLE_DIR / "Lecture_03_Optimization_and_Regularization.srt"
    content = """1
00:00:05,000 --> 00:00:32,000
Welcome back everyone. Today we are diving into advanced optimization algorithms, specifically focusing on the Adam optimizer and why it dominates modern deep learning training.

2
00:00:33,000 --> 00:01:05,000
Standard Stochastic Gradient Descent (SGD) uses a single uniform learning rate across all parameters. In contrast, Adam stands for Adaptive Moment Estimation. It maintains two running tracking vectors for each individual parameter weight.

3
00:01:06,000 --> 00:01:45,000
The first moment is the exponential moving average of past gradients, which acts like momentum. The second moment is the exponential moving average of squared past gradients, which scales the learning rate inversely by the variance of the gradient.

4
00:01:46,000 --> 00:02:25,000
Now let us talk about regularization. Overfitting occurs when our model memorizes training noise rather than generalizable patterns. The most common defense in deep networks is Dropout. During training, Dropout randomly zeros out a fraction of neuron activations, forcing the remaining network to learn redundant and robust representations.

5
00:02:26,000 --> 00:02:58,000
Remember that during inference or testing, Dropout is turned off and the weights are scaled appropriately. Batch Normalization also provides an inherent regularizing effect by normalizing activations across the mini-batch dimension.
"""
    srt_path.write_text(content, encoding="utf-8")
    print(f"Created sample SRT: {srt_path}")

def create_sample_md():
    md_path = SAMPLE_DIR / "Notes_Evaluation_Metrics_and_Losses.md"
    content = """# Machine Learning Evaluation & Loss Functions

## Classification Loss Functions
Binary Cross-Entropy (Log Loss) measures performance where output probabilities range between 0 and 1.
For multi-class classification problems with C mutually exclusive classes, Categorical Cross-Entropy is used together with a Softmax output layer.

## Key Evaluation Metrics
- **Accuracy**: Proportion of total correct predictions. Deceptive when handling severely imbalanced datasets.
- **Precision**: True Positives / (True Positives + False Positives). Critical when the cost of False Positives is severe (e.g. spam detection).
- **Recall (Sensitivity)**: True Positives / (True Positives + False Negatives). Essential in medical diagnostics where missing a positive case is catastrophic.
- **F1 Score**: Harmonic mean of Precision and Recall: 2 * (Precision * Recall) / (Precision + Recall).

## Learning Rate Schedules
- **Cosine Annealing**: Decays the learning rate following a cosine curve down to a minimum threshold.
- **Warmup Phase**: Linearly increases learning rate for the initial 5-10% of training steps to stabilize optimization.
"""
    md_path.write_text(content, encoding="utf-8")
    print(f"Created sample Markdown: {md_path}")

if __name__ == "__main__":
    create_sample_pdf()
    create_sample_pptx()
    create_sample_srt()
    create_sample_md()
